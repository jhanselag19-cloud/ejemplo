from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Q, F, Count, Sum, ProtectedError
from django.http import HttpResponse
from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.views.generic import (
    ListView, DetailView, CreateView, UpdateView, DeleteView, TemplateView,
)

import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, timedelta
from decimal import Decimal

from .models import (
    Producto, MovimientoStock, DetalleMovimiento,
    SolicitudCompra, DetalleSolicitudCompra,
    AreaAlmacen, AlertaStock,
    Transferencia, DetalleTransferencia,
    Notificacion,
)
from .forms import (
    ProductoForm, ProductoFiltroForm,
    SolicitudCompraForm, DetalleSolicitudCompraFormSet, FiltroSolicitudCompraForm,
    AreaAlmacenForm,
    TransferenciaForm, DetalleTransferenciaFormSet,
)


# =============================================================================
# HELPER — REGISTRO AUTOMÁTICO DE MOVIMIENTOS
# =============================================================================

def registrar_movimiento_automatico(producto, tipo, usuario, motivo, cantidad=None):
    """Crea un MovimientoStock + DetalleMovimiento automático."""
    if not producto.area:
        print(f"❌ NO SE CREA: producto {producto.sku} sin área")
        return None

    if not usuario or not usuario.is_authenticated:
        print(f"❌ NO SE CREA: usuario no autenticado")
        return None

    if cantidad is None:
        cantidad = producto.stock_actual

    if cantidad <= 0:
        print(f"❌ NO SE CREA: cantidad {cantidad} no es válida")
        return None

    movimiento = MovimientoStock.objects.create(
        tipo=tipo,
        estado=MovimientoStock.ESTADO_CONFIRMADO,
        area=producto.area,
        proveedor=producto.proveedor,
        creado_por=usuario,
        confirmado_por=usuario,
        motivo=motivo,
    )

    DetalleMovimiento.objects.create(
        movimiento=movimiento,
        producto=producto,
        cantidad=abs(cantidad),
        precio_unitario=producto.precio_unitario or 0,
    )

    print(f"✅ MOVIMIENTO CREADO: {movimiento.numero} | producto={producto.sku} | cantidad={cantidad}")
    return movimiento


# =============================================================================
# HOME
# =============================================================================

def home(request):
    """Dashboard principal con contadores y gráfica de movimientos."""
    total_productos = Producto.objects.count()
    bajo_minimo = Producto.objects.filter(
        activo=True,
        stock_actual__lte=F("stock_minimo"),
        stock_actual__gt=0
    ).count()
    sin_stock = Producto.objects.filter(activo=True, stock_actual=0).count()

    valor_inventario = sum(
        (p.stock_actual or 0) * (p.precio_unitario or 0)
        for p in Producto.objects.filter(activo=True)
    )

    hoy = timezone.now().date()
    dias_es = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]

    dias_grafica = []
    entradas_grafica = []
    salidas_grafica = []

    for i in range(6, -1, -1):
        dia = hoy - timedelta(days=i)
        detalles_dia = DetalleMovimiento.objects.filter(
            movimiento__estado=MovimientoStock.ESTADO_CONFIRMADO,
            movimiento__fecha_movimiento__date=dia,
        )
        entradas = sum(d.cantidad for d in detalles_dia if d.es_entrada)
        salidas = sum(d.cantidad for d in detalles_dia if not d.es_entrada)
        dias_grafica.append(dias_es[dia.weekday()])
        entradas_grafica.append(float(entradas))
        salidas_grafica.append(float(salidas))

    productos_recientes = Producto.objects.filter(activo=True).order_by("-creado_en")[:5]
    productos_atencion = Producto.objects.filter(
        activo=True,
        stock_actual__lte=F("stock_minimo")
    ).order_by("stock_actual")[:5]

    contexto = {
        "total_productos": total_productos,
        "bajo_minimo": bajo_minimo,
        "sin_stock": sin_stock,
        "valor_inventario": valor_inventario,
        "dias_grafica": dias_grafica,
        "entradas_grafica": entradas_grafica,
        "salidas_grafica": salidas_grafica,
        "productos_recientes": productos_recientes,
        "productos_atencion": productos_atencion,
        "total_atencion": bajo_minimo + sin_stock,
    }
    return render(request, "almacen/home.html", contexto)


# =============================================================================
# PRODUCTOS — LISTADO CON FILTROS
# =============================================================================

class ProductoListView(ListView):
    model               = Producto
    template_name       = "almacen/producto_list.html"
    context_object_name = "productos"
    paginate_by         = 10

    def get_queryset(self):
        qs = Producto.objects.select_related(
            "area", "categoria", "unidad_medida", "estado_fisico", "proveedor"
        )
        form = ProductoFiltroForm(self.request.GET or None)
        if form.is_valid():
            data = form.cleaned_data
            if data.get("q"):
                q = data["q"]
                qs = qs.filter(
                    Q(sku__icontains=q) | Q(nombre__icontains=q) |
                    Q(descripcion__icontains=q) | Q(numero_serie__icontains=q) |
                    Q(codigo_barras__icontains=q)
                )
            if data.get("area"):
                qs = qs.filter(area=data["area"])
            if data.get("categoria"):
                qs = qs.filter(categoria=data["categoria"])
            if data.get("estado_fisico"):
                qs = qs.filter(estado_fisico=data["estado_fisico"])
            if data.get("proveedor"):
                qs = qs.filter(proveedor=data["proveedor"])

            stock_estado = data.get("stock_estado")
            if stock_estado == "bajo_minimo":
                qs = qs.filter(stock_actual__lte=F("stock_minimo"), stock_actual__gt=0)
            elif stock_estado == "cerca_minimo":
                qs = qs.filter(
                    stock_actual__gt=F("stock_minimo"),
                    stock_actual__lte=F("umbral_alerta"),
                ) | qs.filter(
                    umbral_alerta=0,
                    stock_actual__gt=F("stock_minimo"),
                    stock_actual__lte=F("stock_minimo") * 1.10,
                )
            elif stock_estado == "sin_stock":
                qs = qs.filter(stock_actual=0)
            elif stock_estado == "ok":
                qs = qs.filter(stock_actual__gt=F("stock_minimo"))

            activo = data.get("activo")
            if activo == "1":
                qs = qs.filter(activo=True)
            elif activo == "0":
                qs = qs.filter(activo=False)

        return qs.order_by("nombre")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["filtro_form"] = ProductoFiltroForm(self.request.GET or None)
        ctx["total_filtrados"] = self.get_queryset().count()
        ctx["bajo_minimo"] = Producto.objects.filter(activo=True, stock_actual__lte=F("stock_minimo")).count()
        ctx["sin_stock"]   = Producto.objects.filter(activo=True, stock_actual=0).count()
        return ctx


# =============================================================================
# PRODUCTOS — EXPORTAR A EXCEL
# =============================================================================

def exportar_productos_excel(request):
    view = ProductoListView()
    view.request = request
    view.kwargs = {}
    view.args = ()
    queryset = view.get_queryset()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario"

    encabezados = [
        "SKU", "Nombre", "Área", "Categoría", "Estado Físico",
        "Stock Actual", "Stock Mínimo", "Precio Unitario",
        "Unidad de Medida", "Proveedor", "Nº Serie", "Código de Barras", "Activo"
    ]
    ws.append(encabezados)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for p in queryset:
        ws.append([
            p.sku, p.nombre,
            p.area.nombre if p.area else "—",
            p.categoria.nombre if p.categoria else "—",
            p.estado_fisico.nombre if p.estado_fisico else "—",
            p.stock_actual, p.stock_minimo,
            float(p.precio_unitario) if p.precio_unitario else 0,
            p.unidad_medida.nombre if p.unidad_medida else "—",
            p.proveedor.nombre if p.proveedor else "—",
            p.numero_serie or "—", p.codigo_barras or "—",
            "Sí" if p.activo else "No",
        ])

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 3

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    fecha = datetime.now().strftime("%Y-%m-%d")
    response["Content-Disposition"] = f'attachment; filename="inventario_{fecha}.xlsx"'
    wb.save(response)
    return response


# =============================================================================
# PRODUCTOS — CRUD
# =============================================================================

class ProductoDetailView(DetailView):
    model               = Producto
    template_name       = "almacen/producto_detail.html"
    context_object_name = "producto"
    slug_field          = "codigo_qr"
    slug_url_kwarg      = "codigo_qr"


class ProductoCreateView(CreateView):
    model         = Producto
    form_class    = ProductoForm
    template_name = "almacen/producto_form.html"

    def form_valid(self, form):
        response = super().form_valid(form)
        if self.object.stock_actual > 0:
            registrar_movimiento_automatico(
                producto=self.object,
                tipo=MovimientoStock.TIPO_ENTRADA_COMPRA,
                usuario=self.request.user,
                motivo=f"Stock inicial del producto {self.object.sku} ({self.object.stock_actual} uds.)",
                cantidad=self.object.stock_actual,
            )
        messages.success(self.request, "Producto creado correctamente.")
        return response

    def get_success_url(self):
        return reverse("almacen:producto_detail", kwargs={"codigo_qr": self.object.codigo_qr})


class ProductoUpdateView(UpdateView):
    model         = Producto
    form_class    = ProductoForm
    template_name = "almacen/producto_form.html"
    slug_field    = "codigo_qr"
    slug_url_kwarg = "codigo_qr"

    def form_valid(self, form):
        stock_anterior = Producto.objects.get(pk=self.object.pk).stock_actual
        response = super().form_valid(form)
        stock_nuevo = self.object.stock_actual

        if stock_nuevo != stock_anterior:
            diferencia = stock_nuevo - stock_anterior
            tipo = (
                MovimientoStock.TIPO_AJUSTE_POSITIVO
                if diferencia > 0
                else MovimientoStock.TIPO_AJUSTE_NEGATIVO
            )
            signo = "+" if diferencia > 0 else ""
            registrar_movimiento_automatico(
                producto=self.object,
                tipo=tipo,
                usuario=self.request.user,
                motivo=f"Ajuste de stock: {stock_anterior} → {stock_nuevo} ({signo}{diferencia})",
                cantidad=abs(diferencia),
            )
        messages.success(self.request, "Producto actualizado.")
        return response

    def get_success_url(self):
        return reverse("almacen:producto_detail", kwargs={"codigo_qr": self.object.codigo_qr})


class ProductoDeleteView(DeleteView):
    model               = Producto
    template_name       = "almacen/producto_confirm_delete.html"
    context_object_name = "producto"
    slug_field          = "codigo_qr"
    slug_url_kwarg      = "codigo_qr"
    success_url         = reverse_lazy("almacen:producto_list")

    def form_valid(self, form):
        producto = self.get_object()
        stock_al_eliminar = producto.stock_actual

        if stock_al_eliminar > 0:
            registrar_movimiento_automatico(
                producto=producto,
                tipo=MovimientoStock.TIPO_AJUSTE_NEGATIVO,
                usuario=self.request.user,
                motivo=(
                    f"Eliminación del producto {producto.sku} — {producto.nombre} "
                    f"con {stock_al_eliminar} uds. en stock"
                ),
                cantidad=stock_al_eliminar,
            )

        sku = producto.sku
        nombre = producto.nombre

        try:
            messages.success(
                self.request,
                f"Producto «{sku} — {nombre}» eliminado. "
                f"Se registró un ajuste negativo de {stock_al_eliminar} uds. en el historial."
            )
            return super().form_valid(form)
        except ProtectedError:
            messages.error(
                self.request,
                f"No se puede eliminar el producto «{sku}» porque algún registro "
                f"todavía lo referencia con PROTECT."
            )
            return redirect("almacen:producto_detail", codigo_qr=producto.codigo_qr)


# =============================================================================
# MOVIMIENTOS
# =============================================================================

class MovimientoListView(ListView):
    model               = MovimientoStock
    template_name       = "almacen/movimiento_list.html"
    context_object_name = "movimientos"
    paginate_by         = 10

    def get_queryset(self):
        qs = MovimientoStock.objects.select_related(
            "area", "proveedor", "departamento", "centro_costo",
            "creado_por", "confirmado_por"
        ).prefetch_related("detalles__producto")

        tipo = self.request.GET.get("tipo")
        if tipo:
            qs = qs.filter(tipo=tipo)

        return qs.order_by("-fecha_movimiento", "-id")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["total_movimientos"] = MovimientoStock.objects.count()
        return ctx


class MovimientoDetailView(DetailView):
    model               = MovimientoStock
    template_name       = "almacen/movimiento_detail.html"
    context_object_name = "movimiento"
    slug_field          = "numero"
    slug_url_kwarg      = "numero"

    def get_queryset(self):
        return MovimientoStock.objects.select_related(
            "area", "proveedor", "departamento", "centro_costo",
            "creado_por", "confirmado_por"
        ).prefetch_related("detalles__producto")


# =============================================================================
# ÁREAS DE ALMACENAJE
# =============================================================================

class AreaAlmacenListView(LoginRequiredMixin, ListView):
    model               = AreaAlmacen
    template_name       = "almacen/area_list.html"
    context_object_name = "areas"
    paginate_by         = 12

    def get_queryset(self):
        qs = AreaAlmacen.objects.annotate(
            alertas_count=Count(
                'productos__alertas',
                filter=Q(productos__alertas__estado=AlertaStock.ESTADO_PENDIENTE),
                distinct=True
            )
        ).order_by("nombre")

        q = self.request.GET.get("q")
        if q:
            qs = qs.filter(Q(nombre__icontains=q) | Q(codigo__icontains=q))

        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["total_areas"] = AreaAlmacen.objects.count()
        ctx["areas_activas"] = AreaAlmacen.objects.filter(activa=True).count()
        return ctx


class AreaAlmacenCreateView(LoginRequiredMixin, CreateView):
    model         = AreaAlmacen
    form_class    = AreaAlmacenForm
    template_name = "almacen/area_form.html"

    def form_valid(self, form):
        messages.success(self.request, "Área de almacenaje creada correctamente.")
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("almacen:area_list")


class AreaAlmacenUpdateView(LoginRequiredMixin, UpdateView):
    model          = AreaAlmacen
    form_class     = AreaAlmacenForm
    template_name  = "almacen/area_form.html"

    def form_valid(self, form):
        messages.success(self.request, "Área de almacenaje actualizada.")
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("almacen:area_list")


class AreaAlmacenDeleteView(LoginRequiredMixin, DeleteView):
    model               = AreaAlmacen
    template_name       = "almacen/area_confirm_delete.html"
    context_object_name = "area"
    success_url         = reverse_lazy("almacen:area_list")

    def form_valid(self, form):
        if self.object.productos.exists():
            messages.error(
                self.request,
                f"No se puede eliminar el área '{self.object.nombre}' porque tiene "
                f"{self.object.productos.count()} producto(s) asociado(s)."
            )
            return redirect("almacen:area_list")

        messages.success(self.request, "Área de almacenaje eliminada.")
        return super().form_valid(form)


# =============================================================================
# REPORTES
# =============================================================================

class ReportesView(LoginRequiredMixin, TemplateView):
    template_name = "almacen/reportes.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        tab = self.request.GET.get("tab", "valoracion")
        ctx["tab"] = tab

        # ✅ Paleta de colores compartida (fuera de los if)
        PALETA_COLORES = [
            "#1B2A4A", "#E30613", "#0D9488", "#7C3AED", "#EA580C",
            "#65A30D", "#0EA5E9", "#8B5CF6", "#F59E0B", "#EF4444",
            "#14B8A6", "#6366F1", "#F97316", "#84CC16", "#06B6D4",
        ]

        # =====================================================================
        # TAB VALORACIÓN
        # =====================================================================
        if tab == "valoracion":
            areas = AreaAlmacen.objects.filter(activa=True).prefetch_related("productos")
            filas = []
            total_global = Decimal("0.00")
            total_productos = 0

            # Primera pasada: recolectar datos básicos
            datos_areas = []
            for idx, area in enumerate(areas):
                productos_area = area.productos.filter(activo=True)
                num_productos = productos_area.count()
                valor_area = sum(
                    (p.stock_actual or 0) * (p.precio_unitario or Decimal("0"))
                    for p in productos_area
                )
                datos_areas.append({
                    "nombre": area.nombre,
                    "codigo": area.codigo,
                    "productos": num_productos,
                    "valoracion": valor_area,
                    "color": PALETA_COLORES[idx % len(PALETA_COLORES)],
                })
                total_global += valor_area
                total_productos += num_productos

            # Segunda pasada: calcular porcentajes y preparar arrays
            areas_labels = []
            areas_valores = []
            areas_colores = []

            for dato in datos_areas:
                porcentaje = float((dato["valoracion"] / total_global) * 100) if total_global > 0 else 0
                dato["porcentaje"] = porcentaje
                filas.append(dato)

                areas_labels.append(dato["nombre"])
                areas_valores.append(float(dato["valoracion"]))
                areas_colores.append(dato["color"])

            # Área top y bottom (solo entre las que tienen valor > 0)
            areas_con_valor = [f for f in filas if f["valoracion"] > 0]
            areas_ordenadas = sorted(areas_con_valor, key=lambda x: x["valoracion"], reverse=True)

            area_top = areas_ordenadas[0] if areas_ordenadas else None
            area_bottom = areas_ordenadas[-1] if len(areas_ordenadas) > 1 else None

            ctx["filas_valoracion"] = filas
            ctx["total_global"] = total_global
            ctx["total_productos_global"] = total_productos
            ctx["total_areas_global"] = len([f for f in filas if f["productos"] > 0])
            ctx["area_top"] = area_top
            ctx["area_bottom"] = area_bottom

            # JSON para las gráficas
            ctx["areas_labels"]  = json.dumps(areas_labels)
            ctx["areas_valores"] = json.dumps(areas_valores)
            ctx["areas_colores"] = json.dumps(areas_colores)

        # =====================================================================
        # TAB MOVIMIENTOS
        # =====================================================================
        elif tab == "movimientos":
            fecha_hasta_str = self.request.GET.get("hasta", "")
            fecha_desde_str = self.request.GET.get("desde", "")

            try:
                fecha_hasta = datetime.strptime(fecha_hasta_str, "%Y-%m-%d").date() if fecha_hasta_str else timezone.now().date()
                fecha_desde = datetime.strptime(fecha_desde_str, "%Y-%m-%d").date() if fecha_desde_str else (fecha_hasta - timedelta(days=30))
            except ValueError:
                fecha_hasta = timezone.now().date()
                fecha_desde = fecha_hasta - timedelta(days=30)

            ctx["fecha_desde"] = fecha_desde.strftime("%Y-%m-%d")
            ctx["fecha_hasta"] = fecha_hasta.strftime("%Y-%m-%d")

            detalles = DetalleMovimiento.objects.filter(
                movimiento__estado=MovimientoStock.ESTADO_CONFIRMADO,
                movimiento__fecha_movimiento__date__gte=fecha_desde,
                movimiento__fecha_movimiento__date__lte=fecha_hasta,
            ).select_related("movimiento", "producto", "producto__area")

            areas = AreaAlmacen.objects.filter(activa=True)
            filas = []
            total_movimientos = 0
            total_entradas = 0
            total_salidas = 0

            for idx, area in enumerate(areas):
                detalles_area = [d for d in detalles if d.producto and d.producto.area_id == area.id]
                # ✅ Convertir a float para evitar el error de JSON
                entradas = float(sum(d.cantidad for d in detalles_area if d.es_entrada))
                salidas = float(sum(d.cantidad for d in detalles_area if not d.es_entrada))
                productos_unicos = len(set(d.producto_id for d in detalles_area))
                total_mov_area = len(detalles_area)

                total_movimientos += total_mov_area
                total_entradas += entradas
                total_salidas += salidas

                filas.append({
                    "nombre": area.nombre,
                    "codigo": area.codigo,
                    "productos": productos_unicos,
                    "entradas": entradas,
                    "salidas": salidas,
                    "total_mov": total_mov_area,
                    "color": PALETA_COLORES[idx % len(PALETA_COLORES)],
                })

            # Calcular porcentajes
            for fila in filas:
                fila["porcentaje"] = (
                    (fila["total_mov"] / total_movimientos * 100) if total_movimientos > 0 else 0
                )

            # Áreas más y menos activas (solo las que tienen movimientos)
            filas_con_mov = [f for f in filas if f["total_mov"] > 0]
            filas_ordenadas = sorted(filas_con_mov, key=lambda x: x["total_mov"], reverse=True)

            area_mas_activa = filas_ordenadas[0] if filas_ordenadas else None
            area_menos_activa = filas_ordenadas[-1] if len(filas_ordenadas) > 1 else None

            ctx["filas_movimientos"] = filas
            ctx["total_movimientos_global"] = total_movimientos
            ctx["total_entradas_global"] = total_entradas
            ctx["total_salidas_global"] = total_salidas
            ctx["area_mas_activa"] = area_mas_activa
            ctx["area_menos_activa"] = area_menos_activa

            # JSON para las gráficas (todos los valores ya son float o int)
            ctx["mov_areas_labels"]   = json.dumps([f["nombre"] for f in filas])
            ctx["mov_areas_entradas"] = json.dumps([f["entradas"] for f in filas])
            ctx["mov_areas_salidas"]  = json.dumps([f["salidas"] for f in filas])
            ctx["mov_areas_totales"]  = json.dumps([f["total_mov"] for f in filas])
            ctx["mov_areas_colores"]  = json.dumps([f["color"] for f in filas])

        # =====================================================================
        # TAB AUDITORÍA
        # =====================================================================
        elif tab == "auditoria":
            productos = Producto.objects.filter(activo=True).select_related(
                "area", "unidad_medida"
            ).order_by("area__nombre", "nombre")
            filas = []
            for p in productos:
                filas.append({
                    "sku": p.sku,
                    "nombre": p.nombre,
                    "stock_sistema": p.stock_actual,
                    "unidad": p.unidad_medida.codigo if p.unidad_medida else "",
                })
            ctx["filas_auditoria"] = filas

        return ctx


# =============================================================================
# TRANSFERENCIAS
# =============================================================================

class TransferenciaListView(LoginRequiredMixin, ListView):
    model               = Transferencia
    template_name       = "almacen/transferencia_list.html"
    context_object_name = "transferencias"
    paginate_by         = 10

    def get_queryset(self):
        qs = Transferencia.objects.select_related(
            "area_origen", "area_destino", "creado_por", "confirmado_por"
        ).prefetch_related("detalles__producto")

        q = self.request.GET.get("q")
        if q:
            qs = qs.filter(
                Q(numero__icontains=q) |
                Q(motivo__icontains=q) |
                Q(area_origen__nombre__icontains=q) |
                Q(area_destino__nombre__icontains=q)
            )

        estado = self.request.GET.get("estado")
        if estado:
            qs = qs.filter(estado=estado)

        return qs.order_by("-fecha_transferencia", "-id")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["total_transferencias"] = Transferencia.objects.count()
        ctx["borradores"]   = Transferencia.objects.filter(estado=Transferencia.ESTADO_BORRADOR).count()
        ctx["confirmadas"]  = Transferencia.objects.filter(estado=Transferencia.ESTADO_CONFIRMADA).count()
        ctx["anuladas"]     = Transferencia.objects.filter(estado=Transferencia.ESTADO_ANULADA).count()
        return ctx


class TransferenciaDetailView(LoginRequiredMixin, DetailView):
    model               = Transferencia
    template_name       = "almacen/transferencia_detail.html"
    context_object_name = "transferencia"
    slug_field          = "numero"
    slug_url_kwarg      = "numero"

    def get_queryset(self):
        return Transferencia.objects.select_related(
            "area_origen", "area_destino", "creado_por", "confirmado_por",
            "movimiento_generado"
        ).prefetch_related("detalles__producto")


class TransferenciaCreateView(LoginRequiredMixin, CreateView):
    model         = Transferencia
    form_class    = TransferenciaForm
    template_name = "almacen/transferencia_form.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        if self.request.POST:
            ctx["detalles_formset"] = DetalleTransferenciaFormSet(self.request.POST)
        else:
            ctx["detalles_formset"] = DetalleTransferenciaFormSet()
        return ctx

    def form_valid(self, form):
        ctx = self.get_context_data()
        formset = ctx["detalles_formset"]

        if not formset.is_valid():
            return self.form_invalid(form)

        form.instance.creado_por = self.request.user
        self.object = form.save()
        formset.instance = self.object
        formset.save()

        from django.contrib.auth import get_user_model
        User = get_user_model()
        admins = User.objects.filter(is_staff=True)

        for admin in admins:
            Notificacion.objects.create(
                usuario=admin,
                titulo=f"Nueva transferencia {self.object.numero}",
                mensaje=(
                    f"{self.request.user.get_full_name() or self.request.user.username} "
                    f"creó una transferencia desde {self.object.area_origen.nombre} "
                    f"hacia {self.object.area_destino.nombre}."
                ),
                tipo=Notificacion.TIPO_TRANSFERENCIA,
                url=reverse("almacen:transferencia_detail", kwargs={"numero": self.object.numero}),
            )

        messages.success(self.request, f"Transferencia {self.object.numero} creada correctamente.")
        return redirect("almacen:transferencia_detail", numero=self.object.numero)


class TransferenciaUpdateView(LoginRequiredMixin, UpdateView):
    model         = Transferencia
    form_class    = TransferenciaForm
    template_name = "almacen/transferencia_form.html"
    slug_field    = "numero"
    slug_url_kwarg = "numero"

    def get_queryset(self):
        return Transferencia.objects.filter(estado=Transferencia.ESTADO_BORRADOR)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        if self.request.POST:
            ctx["detalles_formset"] = DetalleTransferenciaFormSet(self.request.POST, instance=self.object)
        else:
            ctx["detalles_formset"] = DetalleTransferenciaFormSet(instance=self.object)
        return ctx

    def form_valid(self, form):
        ctx = self.get_context_data()
        formset = ctx["detalles_formset"]

        if not formset.is_valid():
            return self.form_invalid(form)

        self.object = form.save()
        formset.instance = self.object
        formset.save()

        messages.success(self.request, f"Transferencia {self.object.numero} actualizada.")
        return redirect("almacen:transferencia_detail", numero=self.object.numero)


class TransferenciaDeleteView(LoginRequiredMixin, DeleteView):
    model               = Transferencia
    template_name       = "almacen/transferencia_confirm_delete.html"
    context_object_name = "transferencia"
    slug_field          = "numero"
    slug_url_kwarg      = "numero"
    success_url         = reverse_lazy("almacen:transferencia_list")

    def get_queryset(self):
        return Transferencia.objects.filter(estado=Transferencia.ESTADO_BORRADOR)

    def form_valid(self, form):
        messages.success(self.request, "Transferencia eliminada.")
        return super().form_valid(form)


@login_required
def transferencia_confirmar(request, numero):
    transferencia = get_object_or_404(Transferencia, numero=numero)
    try:
        transferencia.confirmar(request.user)
        messages.success(
            request,
            f"Transferencia {transferencia.numero} confirmada. "
            f"Se generó el movimiento {transferencia.movimiento_generado.numero}."
        )
    except ValueError as e:
        messages.error(request, str(e))
    return redirect("almacen:transferencia_detail", numero=transferencia.numero)


@login_required
def transferencia_anular(request, numero):
    transferencia = get_object_or_404(Transferencia, numero=numero)
    try:
        transferencia.anular(request.user)
        messages.success(request, f"Transferencia {transferencia.numero} anulada.")
    except ValueError as e:
        messages.error(request, str(e))
    return redirect("almacen:transferencia_detail", numero=transferencia.numero)


# =============================================================================
# MOVIMIENTOS — EXPORTAR A EXCEL
# =============================================================================

def exportar_movimientos_excel(request):
    qs = MovimientoStock.objects.select_related(
        "area", "proveedor", "departamento", "centro_costo",
        "creado_por", "confirmado_por"
    ).prefetch_related("detalles__producto")

    tipo = request.GET.get("tipo")
    if tipo:
        qs = qs.filter(tipo=tipo)

    qs = qs.order_by("-fecha_movimiento", "-id")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Movimientos"

    encabezados = [
        "Número", "Fecha", "Tipo", "Estado", "Área", "Producto", "SKU",
        "Cantidad", "Precio Unitario", "Subtotal",
        "Solicitante", "Confirmado por", "Motivo", "Observación"
    ]
    ws.append(encabezados)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for m in qs:
        detalles = m.detalles.all()
        if detalles:
            for det in detalles:
                ws.append([
                    m.numero,
                    m.fecha_movimiento.strftime("%d/%m/%Y %H:%M"),
                    m.get_tipo_display(),
                    m.get_estado_display(),
                    m.area.nombre if m.area else "—",
                    det.producto.nombre if det.producto else "—",
                    det.producto.sku if det.producto else "—",
                    float(det.cantidad),
                    float(det.precio_unitario),
                    float(det.subtotal),
                    (m.creado_por.get_full_name() or m.creado_por.username) if m.creado_por else "—",
                    (m.confirmado_por.get_full_name() or m.confirmado_por.username) if m.confirmado_por else "—",
                    m.motivo or "—",
                    m.observacion or "—",
                ])
        else:
            ws.append([
                m.numero,
                m.fecha_movimiento.strftime("%d/%m/%Y %H:%M"),
                m.get_tipo_display(),
                m.get_estado_display(),
                m.area.nombre if m.area else "—",
                "—", "—", 0, 0, 0,
                (m.creado_por.username) if m.creado_por else "—",
                (m.confirmado_por.username) if m.confirmado_por else "—",
                m.motivo or "—",
                m.observacion or "—",
            ])

    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = max_length + 3

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    fecha = datetime.now().strftime("%Y-%m-%d")
    response["Content-Disposition"] = f'attachment; filename="movimientos_{fecha}.xlsx"'
    wb.save(response)
    return response


# =============================================================================
# COMPRAS — SOLICITUDES DE COMPRA
# =============================================================================

class SolicitudCompraListView(LoginRequiredMixin, ListView):
    model               = SolicitudCompra
    template_name       = "almacen/solicitud_compra_list.html"
    context_object_name = "solicitudes"
    paginate_by         = 10

    def get_queryset(self):
        qs = SolicitudCompra.objects.select_related(
            "area_solicitante", "departamento", "centro_costo",
            "proveedor", "creado_por", "aprobado_por"
        ).prefetch_related("detalles__producto")

        q = self.request.GET.get("q")
        if q:
            qs = qs.filter(Q(numero__icontains=q) | Q(motivo__icontains=q))

        estado = self.request.GET.get("estado")
        if estado:
            qs = qs.filter(estado=estado)

        area = self.request.GET.get("area")
        if area:
            qs = qs.filter(area_solicitante_id=area)

        return qs.order_by("-fecha_solicitud", "-id")

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["filtro_form"] = FiltroSolicitudCompraForm(self.request.GET or None)
        ctx["total_solicitudes"] = SolicitudCompra.objects.count()
        ctx["pendientes"]        = SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_PENDIENTE).count()
        ctx["aprobadas"]         = SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_APROBADA).count()
        ctx["rechazadas"]        = SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_RECHAZADA).count()
        ctx["compradas"]         = SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_COMPRADA).count()
        return ctx


class SolicitudCompraDetailView(LoginRequiredMixin, DetailView):
    model               = SolicitudCompra
    template_name       = "almacen/solicitud_compra_detail.html"
    context_object_name = "solicitud"
    slug_field          = "numero"
    slug_url_kwarg      = "numero"

    def get_queryset(self):
        return SolicitudCompra.objects.select_related(
            "area_solicitante", "departamento", "centro_costo",
            "proveedor", "creado_por", "aprobado_por", "movimiento_generado"
        ).prefetch_related("detalles__producto")


class SolicitudCompraCreateView(LoginRequiredMixin, CreateView):
    model         = SolicitudCompra
    form_class    = SolicitudCompraForm
    template_name = "almacen/solicitud_compra_form.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        if self.request.POST:
            ctx["detalles_formset"] = DetalleSolicitudCompraFormSet(self.request.POST)
        else:
            ctx["detalles_formset"] = DetalleSolicitudCompraFormSet()
        return ctx

    def form_valid(self, form):
        ctx = self.get_context_data()
        formset = ctx["detalles_formset"]

        if not formset.is_valid():
            return self.form_invalid(form)

        form.instance.creado_por = self.request.user
        self.object = form.save()
        formset.instance = self.object
        formset.save()

        from django.contrib.auth import get_user_model
        User = get_user_model()
        admins = User.objects.filter(is_staff=True)

        for admin in admins:
            Notificacion.objects.create(
                usuario=admin,
                titulo=f"Nueva solicitud de compra {self.object.numero}",
                mensaje=(
                    f"{self.request.user.get_full_name() or self.request.user.username} "
                    f"creó una solicitud para {self.object.area_solicitante.nombre}. "
                    f"Motivo: {self.object.motivo}"
                ),
                tipo=Notificacion.TIPO_SOLICITUD_COMPRA,
                url=reverse("almacen:solicitud_compra_detail", kwargs={"numero": self.object.numero}),
            )

        messages.success(self.request, f"Solicitud {self.object.numero} creada correctamente.")
        return redirect("almacen:solicitud_compra_detail", numero=self.object.numero)


class SolicitudCompraUpdateView(LoginRequiredMixin, UpdateView):
    model         = SolicitudCompra
    form_class    = SolicitudCompraForm
    template_name = "almacen/solicitud_compra_form.html"
    slug_field    = "numero"
    slug_url_kwarg = "numero"

    def get_queryset(self):
        return SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_PENDIENTE)

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        if self.request.POST:
            ctx["detalles_formset"] = DetalleSolicitudCompraFormSet(self.request.POST, instance=self.object)
        else:
            ctx["detalles_formset"] = DetalleSolicitudCompraFormSet(instance=self.object)
        return ctx

    def form_valid(self, form):
        ctx = self.get_context_data()
        formset = ctx["detalles_formset"]

        if not formset.is_valid():
            return self.form_invalid(form)

        self.object = form.save()
        formset.instance = self.object
        formset.save()

        messages.success(self.request, f"Solicitud {self.object.numero} actualizada.")
        return redirect("almacen:solicitud_compra_detail", numero=self.object.numero)


class SolicitudCompraDeleteView(LoginRequiredMixin, DeleteView):
    model               = SolicitudCompra
    template_name       = "almacen/solicitud_compra_confirm_delete.html"
    context_object_name = "solicitud"
    slug_field          = "numero"
    slug_url_kwarg      = "numero"
    success_url         = reverse_lazy("almacen:solicitud_compra_list")

    def get_queryset(self):
        return SolicitudCompra.objects.filter(estado=SolicitudCompra.ESTADO_PENDIENTE)

    def form_valid(self, form):
        messages.success(self.request, "Solicitud eliminada.")
        return super().form_valid(form)


# =============================================================================
# COMPRAS — ACCIONES DE ESTADO
# =============================================================================

@login_required
def solicitud_compra_aprobar(request, numero):
    solicitud = get_object_or_404(SolicitudCompra, numero=numero)
    try:
        solicitud.aprobar(request.user)
        messages.success(request, f"Solicitud {solicitud.numero} aprobada.")
    except ValueError as e:
        messages.error(request, str(e))
    return redirect("almacen:solicitud_compra_detail", numero=solicitud.numero)


@login_required
def solicitud_compra_rechazar(request, numero):
    solicitud = get_object_or_404(SolicitudCompra, numero=numero)
    if request.method == "POST":
        motivo = request.POST.get("motivo_rechazo", "")
        try:
            solicitud.rechazar(request.user, motivo)
            messages.success(request, f"Solicitud {solicitud.numero} rechazada.")
        except ValueError as e:
            messages.error(request, str(e))
    return redirect("almacen:solicitud_compra_detail", numero=solicitud.numero)


@login_required
def solicitud_compra_marcar_comprada(request, numero):
    solicitud = get_object_or_404(SolicitudCompra, numero=numero)
    try:
        solicitud.marcar_comprada(request.user)
        messages.success(
            request,
            f"Solicitud {solicitud.numero} marcada como comprada. "
            f"Se generó el movimiento {solicitud.movimiento_generado.numero}."
        )
    except ValueError as e:
        messages.error(request, str(e))
    return redirect("almacen:solicitud_compra_detail", numero=solicitud.numero)


# =============================================================================
# QR
# =============================================================================

def producto_qr(request, codigo_qr):
    producto = get_object_or_404(Producto, codigo_qr=codigo_qr)
    try:
        import qrcode
        from io import BytesIO
        url_ficha = request.build_absolute_uri(
            reverse("almacen:producto_detail", kwargs={"codigo_qr": producto.codigo_qr})
        )
        img = qrcode.make(url_ficha)
        buffer = BytesIO()
        img.save(buffer, format="PNG")
        return HttpResponse(buffer.getvalue(), content_type="image/png")
    except ImportError:
        return HttpResponse(
            "Falta instalar 'qrcode[pil]'. Ejecuta: pip install qrcode[pil]",
            status=500,
        )


def escanear_qr(request):
    codigo = request.GET.get("codigo", "").strip()

    if codigo:
        producto = None

        try:
            import uuid as uuid_lib
            uuid_valido = uuid_lib.UUID(codigo)
            producto = Producto.objects.filter(codigo_qr=uuid_valido).first()
        except (ValueError, AttributeError):
            pass

        if not producto:
            producto = Producto.objects.filter(sku__iexact=codigo).first()

        if not producto:
            producto = Producto.objects.filter(codigo_barras__iexact=codigo).first()

        if producto:
            messages.success(request, f"Producto encontrado: {producto.sku} — {producto.nombre}")
            return redirect("almacen:producto_detail", codigo_qr=producto.codigo_qr)
        else:
            messages.error(request, f"No se encontró ningún producto con el código: {codigo}")

    return render(request, "almacen/escanear.html")


# =============================================================================
# NOTIFICACIONES
# =============================================================================

@login_required
def notificaciones_marcar_leidas(request):
    Notificacion.objects.filter(usuario=request.user, leida=False).update(
        leida=True,
        leida_en=timezone.now(),
    )
    return redirect("almacen:home")


@login_required
def notificacion_marcar_leida(request, pk):
    notificacion = get_object_or_404(Notificacion, pk=pk, usuario=request.user)
    notificacion.marcar_como_leida()
    if notificacion.url:
        return redirect(notificacion.url)
    return redirect("almacen:home")


# =============================================================================
# PRODUCTOS — DESACTIVAR / REACTIVAR
# =============================================================================

@login_required
def producto_desactivar(request, codigo_qr):
    producto = get_object_or_404(Producto, codigo_qr=codigo_qr)
    producto.activo = False
    producto.save(update_fields=["activo", "actualizado_en"])
    messages.success(
        request,
        f"Producto «{producto.sku}» desactivado. Sigue en el historial pero no aparece en el catálogo activo."
    )
    return redirect("almacen:producto_detail", codigo_qr=producto.codigo_qr)


@login_required
def producto_activar(request, codigo_qr):
    producto = get_object_or_404(Producto, codigo_qr=codigo_qr)
    producto.activo = True
    producto.save(update_fields=["activo", "actualizado_en"])
    messages.success(request, f"Producto «{producto.sku}» reactivado.")
    return redirect("almacen:producto_detail", codigo_qr=producto.codigo_qr)